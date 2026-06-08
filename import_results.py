# import os
# import glob
# import pandas as pd
# from openpyxl import load_workbook
# from openpyxl.styles import Font, PatternFill, Alignment
# from openpyxl.utils import get_column_letter

# # ── PATHS — update if different on your machine ──────────────────────────────
# RESULTS_FOLDER = "C:/Users/okech/Models/DigitalTwin/results/"
# EXCEL_FILE     = "C:/Users/okech/Models/DigitalTwin/results/Sprint3_Results.xlsx"
# SUMMARY_CSV    = os.path.join(RESULTS_FOLDER, "summary_results.csv")
# # ─────────────────────────────────────────────────────────────────────────────

# LIGHT_BLUE  = "D5E8F0"
# LIGHT_GREEN = "E2EFDA"
# LIGHT_AMBER = "FFF2CC"
# WHITE       = "FFFFFF"

# def cell_bg(condition):
#     if condition == 0: return LIGHT_BLUE
#     if condition == 1: return LIGHT_GREEN
#     return LIGHT_AMBER

# def style_cell(cell, bg=WHITE, bold=False, center=True, num_fmt=None):
#     cell.fill = PatternFill("solid", start_color=bg)
#     cell.font = Font(name="Arial", bold=bold, size=10)
#     cell.alignment = Alignment(
#         horizontal="center" if center else "left",
#         vertical="center")
#     if num_fmt:
#         cell.number_format = num_fmt

# CONDITION_LABELS = {0: "A0 - Baseline", 1: "A1 - HITL", 2: "A2 - Full Auto"}
# SCENARIO_LABELS  = {0: "S1 - Stable",   1: "S2 - Shock", 2: "S3 - Disruption"}

# # Row mapping in Raw Data sheet
# # Each cell has 5 replication rows + 1 spacer
# # Order: A0S0, A0S1, A0S2, A1S0, A1S1, A1S2, A2S0, A2S1, A2S2
# def get_excel_row(condition, scenario, replication):
#     cell_index = (condition * 3) + scenario
#     start_row  = 4 + (cell_index * 6)  # 5 data rows + 1 spacer = 6 per cell
#     return start_row + (replication - 1)

# # ── STEP 1: Import summary results ───────────────────────────────────────────
# print("=" * 60)
# print("SPRINT 3 RESULTS IMPORTER")
# print("=" * 60)

# if not os.path.exists(SUMMARY_CSV):
#     print(f"ERROR: Cannot find {SUMMARY_CSV}")
#     print("Make sure the path is correct and the file exists.")
#     exit(1)

# if not os.path.exists(EXCEL_FILE):
#     print(f"ERROR: Cannot find {EXCEL_FILE}")
#     print("Make sure Sprint3_Results.xlsx is in the results folder.")
#     exit(1)

# print(f"\nReading: {SUMMARY_CSV}")
# df = pd.read_csv(SUMMARY_CSV)
# print(f"Found {len(df)} rows of simulation results")
# print(df.to_string(index=False))

# print(f"\nOpening: {EXCEL_FILE}")
# wb = load_workbook(EXCEL_FILE)
# ws = wb["Raw Data"]

# imported = 0
# skipped  = 0

# for _, row in df.iterrows():
#     condition   = int(row["Condition"])
#     scenario    = int(row["Scenario"])
#     replication = int(row["Replication"])

#     excel_row = get_excel_row(condition, scenario, replication)
#     bg = cell_bg(condition)

#     # Skip replication 1 — already pre-filled
#     if replication == 1:
#         print(f"  Skipping A{condition}xS{scenario} Rep1 (already pre-filled)")
#         skipped += 1
#         continue

#     # Write values into columns E-K (5-11)
#     values = [
#         row["RetailerBWR"],
#         row["DistributorBWR"],
#         row["ManufacturerBWR"],
#         row["RetailerCost"],
#         row["DistributorCost"],
#         row["ManufacturerCost"],
#         row["TotalCost"],
#     ]
#     num_fmts = ["0.000","0.000","0.000","#,##0.00","#,##0.00","#,##0.00","#,##0.00"]

#     for col_offset, (val, fmt) in enumerate(zip(values, num_fmts)):
#         cell = ws.cell(row=excel_row, column=5 + col_offset, value=round(float(val), 4))
#         style_cell(cell, bg=bg, num_fmt=fmt)

#     # Update notes column
#     notes_cell = ws.cell(row=excel_row, column=12,
#                          value=f"Rep {replication} — imported")
#     style_cell(notes_cell, bg=bg, center=False)

#     print(f"  Imported A{condition}xS{scenario} Rep{replication} → Excel row {excel_row}")
#     imported += 1

# print(f"\nSummary import complete: {imported} rows imported, {skipped} skipped")

# # ── STEP 2: Import time series data ──────────────────────────────────────────
# print("\nImporting time series data...")

# ts_files = glob.glob(os.path.join(RESULTS_FOLDER, "timeseries_*.csv"))
# print(f"Found {len(ts_files)} time series files")

# if ts_files:
#     # Combine all time series into one dataframe
#     all_ts = []
#     for f in sorted(ts_files):
#         fname = os.path.basename(f)
#         # Parse filename: timeseries_A0_S1_R1.csv
#         try:
#             parts = fname.replace("timeseries_","").replace(".csv","").split("_")
#             cond  = int(parts[0].replace("A",""))
#             scen  = int(parts[1].replace("S",""))
#             rep   = int(parts[2].replace("R",""))
#             ts_df = pd.read_csv(f)
#             ts_df["Condition"]   = cond
#             ts_df["Scenario"]    = scen
#             ts_df["Replication"] = rep
#             ts_df["Run"] = f"A{cond}_S{scen}_R{rep}"
#             all_ts.append(ts_df)
#             print(f"  Read: {fname} ({len(ts_df)} rows)")
#         except Exception as e:
#             print(f"  Skipping {fname}: {e}")

#     if all_ts:
#         combined = pd.concat(all_ts, ignore_index=True)

#         # Write to Combined Time Series sheet
#         if "Combined Time Series" in wb.sheetnames:
#             del wb["Combined Time Series"]

#         ws_ts = wb.create_sheet("Combined Time Series")
#         ws_ts.sheet_view.showGridLines = False
#         ws_ts.freeze_panes = "A3"

#         # Title
#         title = ws_ts.cell(row=1, column=1,
#             value="COMBINED TIME SERIES DATA — ALL RUNS")
#         title.font = Font(name="Arial", bold=True, color="FFFFFF", size=12)
#         title.fill = PatternFill("solid", start_color="1A3A5C")
#         title.alignment = Alignment(horizontal="center", vertical="center")
#         ws_ts.merge_cells(start_row=1, start_column=1, end_row=1, end_column=11)
#         ws_ts.row_dimensions[1].height = 30

#         # Headers
#         headers = ["Run","Condition","Scenario","Replication",
#                    "Week","Tier","Inventory","Backlog",
#                    "Order","Shipped","IncomingOrder"]
#         widths  = [14, 10, 10, 12, 8, 14, 12, 10, 10, 10, 14]

#         for i, (h, w) in enumerate(zip(headers, widths), 1):
#             cell = ws_ts.cell(row=2, column=i, value=h)
#             cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
#             cell.fill = PatternFill("solid", start_color="2E75B6")
#             cell.alignment = Alignment(horizontal="center", vertical="center")
#             ws_ts.column_dimensions[get_column_letter(i)].width = w
#         ws_ts.row_dimensions[2].height = 25

#         tier_colours = {
#             "Retailer":     "D5E8F0",
#             "Distributor":  "E2EFDA",
#             "Manufacturer": "FFF2CC"
#         }

#         # Write data
#         col_order = ["Run","Condition","Scenario","Replication",
#                      "Week","Tier","Inventory","Backlog",
#                      "Order","Shipped","IncomingOrder"]

#         for row_idx, (_, data_row) in enumerate(combined[col_order].iterrows(), 3):
#             tier = data_row["Tier"]
#             bg   = tier_colours.get(tier, "FFFFFF")
#             for col_idx, col_name in enumerate(col_order, 1):
#                 val  = data_row[col_name]
#                 cell = ws_ts.cell(row=row_idx, column=col_idx, value=val)
#                 cell.font      = Font(name="Arial", size=9)
#                 cell.fill      = PatternFill("solid", start_color=bg)
#                 cell.alignment = Alignment(horizontal="center", vertical="center")
#                 if col_name in ["Inventory","Backlog","Order","Shipped","IncomingOrder"]:
#                     cell.number_format = "0.0"

#         print(f"\nTime series sheet created: {len(combined)} total rows")

# wb.save(EXCEL_FILE)
# print(f"\nExcel file saved: {EXCEL_FILE}")
# print("\n✓ Import complete. Open Sprint3_Results.xlsx to review your data.")
# print("=" * 60)




import os
import glob
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

RESULTS_FOLDER = "C:/Users/okech/Models/DigitalTwin/results/"
EXCEL_FILE     = "C:/Users/okech/Models/DigitalTwin/results/Sprint3_Results.xlsx"
SUMMARY_CSV    = os.path.join(RESULTS_FOLDER, "summary_results.csv")

LIGHT_BLUE   = "D5E8F0"
LIGHT_GREEN  = "E2EFDA"
LIGHT_AMBER  = "FFF2CC"
LIGHT_PURPLE = "EAD6F5"
LIGHT_PINK   = "F5D6D6"
WHITE        = "FFFFFF"

CONDITION_COLOURS = {
    0: LIGHT_BLUE,
    1: LIGHT_GREEN,
    2: LIGHT_AMBER,
    3: LIGHT_PURPLE,
    4: LIGHT_PINK
}

CONDITION_LABELS = {
    0: "A0 - Baseline",
    1: "A1 - HITL",
    2: "A2 - Full Auto",
    3: "A3 - Info Sharing",
    4: "A4 - LLM + RL"
}

SCENARIO_LABELS = {
    0: "S1 - Stable",
    1: "S2 - Shock",
    2: "S3 - Disruption"
}

def style_cell(cell, bg=WHITE, bold=False, center=True, num_fmt=None):
    cell.fill      = PatternFill("solid", start_color=bg)
    cell.font      = Font(name="Arial", bold=bold, size=10)
    cell.alignment = Alignment(
        horizontal="center" if center else "left",
        vertical="center")
    if num_fmt:
        cell.number_format = num_fmt

def get_excel_row(condition, scenario, replication):
    # 5 conditions x 3 scenarios = 15 cells
    # Each cell: 5 data rows + 1 spacer = 6 rows
    cell_index = (condition * 3) + scenario
    start_row  = 4 + (cell_index * 6)
    return start_row + (replication - 1)

print("=" * 60)
print("SPRINT 3 RESULTS IMPORTER — A0 through A4")
print("=" * 60)

if not os.path.exists(SUMMARY_CSV):
    print(f"ERROR: Cannot find {SUMMARY_CSV}")
    exit(1)

if not os.path.exists(EXCEL_FILE):
    print(f"ERROR: Cannot find {EXCEL_FILE}")
    exit(1)

print(f"\nReading: {SUMMARY_CSV}")
df = pd.read_csv(SUMMARY_CSV)
print(f"Found {len(df)} rows")
print(df.to_string(index=False))

print(f"\nOpening: {EXCEL_FILE}")
wb = load_workbook(EXCEL_FILE)

if "Raw Data" not in wb.sheetnames:
    ws = wb.create_sheet("Raw Data")
    headers = ["Condition", "Scenario", "Replication",
               "RetailerBWR", "DistributorBWR", "ManufacturerBWR",
               "RetailerCost", "DistributorCost", "ManufacturerCost",
               "TotalCost", "Notes"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        style_cell(cell, bg="1A3A5C", bold=True)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
else:
    ws = wb["Raw Data"]

imported = 0
skipped  = 0

for _, row in df.iterrows():
    condition   = int(row["Condition"])
    scenario    = int(row["Scenario"])
    replication = int(row["Replication"])
    bg          = CONDITION_COLOURS.get(condition, WHITE)
    excel_row   = get_excel_row(condition, scenario, replication)

    if replication == 1:
        print(f"  Skipping A{condition}xS{scenario} Rep1 (pre-filled)")
        skipped += 1
        continue

    values   = [
        row["RetailerBWR"], row["DistributorBWR"], row["ManufacturerBWR"],
        row["RetailerCost"], row["DistributorCost"], row["ManufacturerCost"],
        row["TotalCost"]
    ]
    num_fmts = ["0.000","0.000","0.000","#,##0.00","#,##0.00","#,##0.00","#,##0.00"]

    for col_offset, (val, fmt) in enumerate(zip(values, num_fmts)):
        cell = ws.cell(row=excel_row, column=5 + col_offset, value=round(float(val), 4))
        style_cell(cell, bg=bg, num_fmt=fmt)

    label = CONDITION_LABELS.get(condition, f"A{condition}")
    notes = ws.cell(row=excel_row, column=12,
                    value=f"Rep {replication} — {label} imported")
    style_cell(notes, bg=bg, center=False)

    print(f"  Imported {label} x {SCENARIO_LABELS.get(scenario, 'S'+str(scenario))}"
          f" Rep{replication} -> row {excel_row}")
    imported += 1

print(f"\nImport: {imported} rows imported, {skipped} skipped")

# ── Summary sheet ─────────────────────────────────────────────────────────────
if "Summary" not in wb.sheetnames:
    ws_sum = wb.create_sheet("Summary")
else:
    ws_sum = wb["Summary"]
    ws_sum.delete_rows(1, ws_sum.max_row)

title = ws_sum.cell(row=1, column=1, value="BWR SUMMARY — MEAN ACROSS REPLICATIONS")
title.font      = Font(name="Arial", bold=True, color="FFFFFF", size=12)
title.fill      = PatternFill("solid", start_color="1A3A5C")
title.alignment = Alignment(horizontal="center", vertical="center")
ws_sum.merge_cells("A1:I1")
ws_sum.row_dimensions[1].height = 28

headers = ["Condition", "Scenario",
           "Retailer BWR", "Distributor BWR", "Manufacturer BWR",
           "Retailer Cost", "Distributor Cost", "Manufacturer Cost", "Total Cost"]
for col, h in enumerate(headers, 1):
    cell = ws_sum.cell(row=2, column=col, value=h)
    cell.font      = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    cell.fill      = PatternFill("solid", start_color="2E75B6")
    cell.alignment = Alignment(horizontal="center", vertical="center")

for col_letter, width in zip(["A","B","C","D","E","F","G","H","I"],
                              [22, 18, 16, 18, 20, 16, 18, 20, 14]):
    ws_sum.column_dimensions[col_letter].width = width

data_row = 3
for condition in sorted(df["Condition"].unique()):
    for scenario in sorted(df["Scenario"].unique()):
        subset = df[(df["Condition"] == condition) & (df["Scenario"] == scenario)]
        if subset.empty:
            continue
        bg    = CONDITION_COLOURS.get(condition, WHITE)
        means = subset[["RetailerBWR","DistributorBWR","ManufacturerBWR",
                         "RetailerCost","DistributorCost","ManufacturerCost",
                         "TotalCost"]].mean()

        row_vals = [
            CONDITION_LABELS.get(condition, f"A{condition}"),
            SCENARIO_LABELS.get(scenario,   f"S{scenario}"),
            round(means["RetailerBWR"],      3),
            round(means["DistributorBWR"],   3),
            round(means["ManufacturerBWR"],  3),
            round(means["RetailerCost"],     2),
            round(means["DistributorCost"],  2),
            round(means["ManufacturerCost"], 2),
            round(means["TotalCost"],        2),
        ]
        fmts = [None, None, "0.000","0.000","0.000",
                "#,##0.00","#,##0.00","#,##0.00","#,##0.00"]

        for col, (val, fmt) in enumerate(zip(row_vals, fmts), 1):
            cell = ws_sum.cell(row=data_row, column=col, value=val)
            style_cell(cell, bg=bg, num_fmt=fmt)

        data_row += 1

wb.save(EXCEL_FILE)
print(f"\nSaved: {EXCEL_FILE}")
print(f"Summary sheet: {data_row - 3} rows written")
print("=" * 60)